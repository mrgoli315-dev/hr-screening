#!/usr/bin/env python3
"""Rate the resumes listed in pipeline/new/chunk_run_NN.json with the Claude API.
Writes pipeline/new/out_run_NN.json in the same format the Claude-in-chat readers produced,
so `update.py merge` works unchanged.

Env:
  ANTHROPIC_API_KEY  required
  CLAUDE_MODEL       default claude-sonnet-5-5 (claude-haiku-5-5 is cheaper)
  RATE_WORKERS       parallel requests, default 4
  GOOGLE_SA_JSON     see gdrive.py
A resume that fails (network, API error) is left out of the output file: it stays
"waiting" on the page and `update.py new` queues it again on the next run.
"""
import base64, glob, io, json, os, re, sys, time
from concurrent.futures import ThreadPoolExecutor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import gdrive
import anthropic

H = os.path.dirname(os.path.abspath(__file__))
NEW = os.path.join(H, 'new')
MODEL = os.environ.get('CLAUDE_MODEL') or 'claude-sonnet-5-5'
WORKERS = int(os.environ.get('RATE_WORKERS') or 4)
MAX_TEXT = 60000

TPL = open(os.path.join(NEW, 'prompt_template.txt'), encoding='utf-8').read()
RULES = TPL[TPL.index("THE COMPANY'S SALES MODEL"):TPL.index('OUTPUT:')].strip()
SCHEMA = re.search(r'^\{"ph".*\}$', TPL, re.M).group(0)

PROMPT = """You are screening one job applicant for MR.GOLISODA (South Asian Foods Corporation, SAFC), a Goli Soda brand that runs district-level mini-factory + distribution franchises in South/West India. The applicant's resume is attached above (or its text is given above). Rate the applicant.

APPLICANT (from the application form): {item}
src = which application form ("BDE", "BDM" or "MBA"; the MBA form asks for MBA holders - an MBA is an added advantage, not a requirement, and the rating stays about sales fit). form = what they typed in the form (self-reported, use only as a hint).
{note}
{rules}

Reply with ONLY one JSON object (no markdown, no text before or after), with exactly these keys:
{schema}"""

client = anthropic.Anthropic()

DOCX = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
IMAGES = {'image/jpeg', 'image/png', 'image/gif', 'image/webp'}


def docx_text(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(p for p in parts if p.strip())


def resume_blocks(fid):
    """Content blocks for the resume, plus a note for the prompt."""
    if not fid:
        return [], 'NOTE: no resume file was given. Rate from the form only with lv "check" and ask "No resume - confirm experience on call".'
    try:
        m = gdrive.meta(fid)
    except Exception as e:
        return [], 'NOTE: the resume file could not be opened (%s). Record lv "hide", why "could not open".' % str(e)[:80]
    mt = m.get('mimeType', '')
    if mt == 'application/vnd.google-apps.document':
        return [{'type': 'text', 'text': 'RESUME TEXT:\n' + gdrive.export_text(fid)[:MAX_TEXT]}], ''
    if mt.startswith('application/vnd.google-apps'):
        return [], 'NOTE: the uploaded file is a Google %s, not a resume. Record lv "hide", why "not a resume".' % mt.split('.')[-1]
    data = gdrive.download(fid)
    if mt == 'application/pdf':
        return [{'type': 'document', 'source': {'type': 'base64', 'media_type': 'application/pdf',
                                                 'data': base64.b64encode(data).decode()}}], ''
    if mt in IMAGES:
        return [{'type': 'image', 'source': {'type': 'base64', 'media_type': mt,
                                              'data': base64.b64encode(data).decode()}}], ''
    if mt == DOCX:
        try:
            return [{'type': 'text', 'text': 'RESUME TEXT:\n' + docx_text(data)[:MAX_TEXT]}], ''
        except Exception:
            pass
    if mt.startswith('text/'):
        return [{'type': 'text', 'text': 'RESUME TEXT:\n' + data.decode('utf-8', 'replace')[:MAX_TEXT]}], ''
    return [], ('NOTE: the resume file "%s" (%s) cannot be read as text. Rate from the form only with lv "check" '
                'and ask "Resume file unreadable - confirm experience on call".' % (m.get('name', ''), mt))


def parse_json(text):
    s, e = text.find('{'), text.rfind('}')
    if s < 0 or e <= s:
        raise ValueError('no JSON in reply')
    return json.loads(text[s:e + 1])


def rate(item):
    blocks, note = resume_blocks(item.get('fid'))
    info = {k: item.get(k, '') for k in ('src', 'n', 'ph', 'st', 'fd', 'form', 'fid')}
    prompt = PROMPT.format(item=json.dumps(info, ensure_ascii=False), note=note, rules=RULES, schema=SCHEMA)
    last = None
    for attempt in range(3):
        try:
            msg = client.messages.create(model=MODEL, max_tokens=1500,
                                         messages=[{'role': 'user', 'content': blocks + [{'type': 'text', 'text': prompt}]}])
            text = ''.join(b.text for b in msg.content if getattr(b, 'type', '') == 'text')
            o = parse_json(text)
            # keep the form's identity fields exactly as given
            for k in ('ph', 'src', 'fd', 'st', 'fid'):
                o[k] = item.get(k, '')
            if o.get('lv') not in ('strong', 'medium', 'check', 'hide'):
                o['lv'] = 'check'
            if not isinstance(o.get('ro'), list) or not o['ro']:
                o['ro'] = ['BDE']
            if not isinstance(o.get('co'), list):
                o['co'] = []
            o.setdefault('n', item.get('n', ''))
            return o
        except anthropic.BadRequestError as e:      # e.g. password-protected or broken PDF
            if blocks:
                blocks, note = [], 'NOTE: the resume file could not be processed (%s). Rate from the form only with lv "check" and ask "Resume file unreadable - get it re-sent".' % str(e)[:80]
                prompt = PROMPT.format(item=json.dumps(info, ensure_ascii=False), note=note, rules=RULES, schema=SCHEMA)
            last = e
        except Exception as e:
            last = e
            time.sleep(5 * (attempt + 1))
    print('FAILED', item.get('n'), item.get('ph'), str(last)[:200], flush=True)
    return None


def main():
    files = sorted(glob.glob(os.path.join(NEW, 'chunk_run_*.json')))
    total = ok = 0
    levels = {}
    for fn in files:
        items = json.load(open(fn, encoding='utf-8'))
        total += len(items)
        with ThreadPoolExecutor(max_workers=WORKERS) as ex:
            res = list(ex.map(rate, items))
        out = [r for r in res if r]
        ok += len(out)
        for r in out:
            levels[r['lv']] = levels.get(r['lv'], 0) + 1
        json.dump(out, open(fn.replace('chunk_run_', 'out_run_'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        print(os.path.basename(fn), len(out), 'of', len(items), flush=True)
    print(json.dumps({'to_rate': total, 'rated': ok, 'levels': levels}))


if __name__ == '__main__':
    main()
